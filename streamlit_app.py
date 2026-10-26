"""لوحة تحكم نظام تقارير المبيعات (Streamlit). كل المنطق في sales_core.py."""
from __future__ import annotations

import datetime as dt
from typing import Any, List

import pandas as pd
import streamlit as st
from openpyxl.utils import get_column_letter

from sales_core import (
    DEFAULT_CATEGORY_REGEX, REQUIRED_SHEETS, Settings, ValidationError, build_output_workbook,
    build_qasadeer, build_statement, load_source, read_table, target_summary,
)

st.set_page_config(page_title="نظام تقارير المبيعات", page_icon="📊", layout="wide")
st.markdown(
    """<style>
    html, body, [data-testid="stAppViewContainer"], [data-testid="stSidebar"] {direction: rtl; text-align: right;}
    [data-testid="stDataFrame"] {direction: ltr;}
    </style>""",
    unsafe_allow_html=True,
)
PREVIEW_ROWS = 300


@st.cache_resource(show_spinner="جاري قراءة الملف...")
def get_source(file_bytes: bytes):
    return load_source(file_bytes)


def unique_headers(headers: List[Any]) -> List[str]:
    out, seen = [], {}
    for i, h in enumerate(headers):
        label = str(h).strip() if h not in (None, "") else f"({get_column_letter(i + 1)})"
        seen[label] = seen.get(label, 0) + 1
        out.append(label if seen[label] == 1 else f"{label} [{get_column_letter(i + 1)}]")
    return out


def sheet_preview(ws, header: bool = False) -> pd.DataFrame:
    rows = list(ws.iter_rows(values_only=True, max_row=PREVIEW_ROWS + 1))
    if not rows:
        return pd.DataFrame()
    df = pd.DataFrame(rows[1:], columns=unique_headers(list(rows[0]))) if header else pd.DataFrame(rows)
    df = df.dropna(axis=1, how="all").dropna(axis=0, how="all")
    return df.astype(object).where(df.notna(), "").astype(str)  # mixed types -> safe to display


# ------------------------------------------------------------------ sidebar
st.sidebar.header("⚙️ الإعدادات")
uploaded = st.sidebar.file_uploader("ارفع ملف Excel (xlsx)", type=["xlsx", "xlsm"])
with st.sidebar.expander("إعدادات متقدمة"):
    name_header = st.text_input("عمود اسم العميل", "اسم العميل")
    number_header = st.text_input("عمود رقم العميل", "رقم العميل")
    region_header = st.text_input("عمود المدينة", "المدينه")
    category_regex = st.text_input("نمط فئات المبالغ (Regex، المجموعة الأولى = المبلغ)", DEFAULT_CATEGORY_REGEX)
    last_row = st.number_input("آخر صف في كتلة العملاء (0 = تلقائي)", min_value=0, value=0, step=1)
    face_total = st.checkbox("إضافة عمود اجمالي المبلغ (القيمة × العدد)", value=False)
    hide_empty = st.checkbox("إخفاء العملاء بلا أي قيم في العرض فقط", value=True)
if st.sidebar.button("🔄 تحديث التقارير"):
    st.cache_resource.clear()
    st.rerun()

st.title("📊 نظام تقارير المبيعات")
if uploaded is None:
    st.info("ارفع ملف Excel من الشريط الجانبي للبدء. ورقة «Data» هي المصدر الأساسي لكل التقارير.")
    st.stop()

# ------------------------------------------------------------------ load + compute
try:
    src = get_source(uploaded.getvalue())
except ValidationError as exc:
    st.error(str(exc))
    st.stop()

for note in src.notes:
    st.warning(note)

settings = Settings(
    name_header=name_header, number_header=number_header, region_header=region_header,
    category_regex=category_regex, include_face_value_total=face_total, last_row=int(last_row) or None,
)
data_ws = src.wb_values[src.sheet_map["Data"]]
headers, rows = read_table(data_ws)
try:
    qas = build_qasadeer(headers, rows, settings)
except ValidationError as exc:
    st.error(str(exc))
    st.stop()

st.success(f"تمت قراءة الملف: {len(qas.table)} عميلًا و{len(qas.categories)} فئة مبلغ ({'، '.join(qas.categories)}).")
for msg in qas.issues:
    st.warning(msg)

# ------------------------------------------------------------------ KPIs
kpi = target_summary(src.wb_values[src.sheet_map["Target"]]) if "Target" in src.sheet_map else None
cols = st.columns(len(qas.categories) + 1)
cols[0].metric("عدد العملاء", f"{len(qas.table):,}")
for c, lab in zip(cols[1:], qas.categories):
    c.metric(f"إجمالي {lab}", f"{qas.table[lab].sum(skipna=True):,.0f}")
if kpi is not None:
    with st.expander("المستهدف مقابل الفعلي 2026 (كما ورد في ورقة Target)", expanded=False):
        st.dataframe(kpi, use_container_width=True, hide_index=True)
        st.caption("القيم منقولة كما هي من الورقة؛ الأشهر التي تظهر 0 قد تعني عدم توفر بيانات بعد.")

# ------------------------------------------------------------------ statement filters (used by tab + export)
tx_name_col = next((i for i, h in enumerate(headers) if str(h).strip().lower() == "customer - name"), None)
customer_names = sorted({str(r[tx_name_col]).strip() for r in rows if tx_name_col is not None and r[tx_name_col]}) if tx_name_col is not None else []

tabs = st.tabs(REQUIRED_SHEETS)
statement_df = None
statement_filtered = False

for tab, name in zip(tabs, REQUIRED_SHEETS):
    with tab:
        if name == "قصادير":
            q = st.text_input("🔎 بحث عن عميل", key="q_search")
            view = qas.table
            if hide_empty:
                view = view[view[qas.categories].notna().any(axis=1)]
            if q:
                view = view[view[settings.name_header.strip()].astype(str).str.contains(q, case=False, na=False)]
            st.dataframe(view, use_container_width=True, hide_index=True)
            st.bar_chart(pd.Series({l: qas.table[l].sum(skipna=True) for l in qas.categories}))
            st.caption("أعمدة الفئات المتكررة في Data تُجمع في عمود واحد لكل فئة، ولا تظهر أعمدة التاريخ.")
        elif name == "Data":
            st.caption(f"معاينة أول {PREVIEW_ROWS} صفًا من أصل {len(rows)}.")
            st.dataframe(sheet_preview(data_ws, header=True), use_container_width=True)
        elif name == "Statement":
            c1, c2, c3 = st.columns(3)
            cust = c1.selectbox("العميل", ["(الكل)"] + customer_names)
            text = c2.text_input("بحث بالاسم", key="st_search")
            use_dates = c3.checkbox("تحديد فترة")
            d_from = d_to = None
            if use_dates:
                d_from = c3.date_input("من", dt.date(2024, 1, 1))
                d_to = c3.date_input("إلى", dt.date.today())
            try:
                statement_df = build_statement(
                    headers, rows, customer=None if cust == "(الكل)" else cust, search=text or None,
                    date_from=d_from, date_to=d_to,
                )
                statement_filtered = cust != "(الكل)" or bool(text) or use_dates
                st.write(f"عدد الصفوف: {len(statement_df):,}")
                st.dataframe(statement_df.astype(object).where(statement_df.notna(), ""), use_container_width=True, hide_index=True)
            except ValidationError as exc:
                st.error(str(exc))
        elif name in src.sheet_map:
            st.dataframe(sheet_preview(src.wb_values[src.sheet_map[name]]), use_container_width=True, hide_index=True)
            st.caption("القيم كما هي محفوظة في الملف (معاينة).")
        else:
            st.warning(f"الورقة «{name}» غير موجودة في الملف المرفوع.")

# ------------------------------------------------------------------ export
st.divider()
st.subheader("⬇️ تنزيل الملف الناتج")
try:
    out_bytes, out_notes = build_output_workbook(src, qas, statement_df, force_statement=statement_filtered)
    for n in out_notes:
        st.info(n)
    st.download_button(
        "تنزيل Total_Sales_Output.xlsx (الأوراق السبع)", out_bytes, file_name="Total_Sales_Output.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    st.caption(
        "الملف يحتفظ بكل الأوراق الأصلية ومعادلاتها، ويعيد بناء «قصادير» من Data. "
        "المعادلات تُحسب عند فتحه في Excel."
    )
except Exception as exc:  # noqa: BLE001
    st.error(f"تعذّر إنشاء الملف الناتج: {exc}")
