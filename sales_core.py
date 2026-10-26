"""Core logic for the sales reporting system (no Streamlit dependency).

Responsibilities, kept separate on purpose:
  1. reading   -> load_source(), read_table()
  2. computing -> build_qasadeer(), build_statement(), target_summary()
  3. writing   -> build_output_workbook()
"""
from __future__ import annotations

import datetime as dt
import io
import re
import unicodedata
from dataclasses import dataclass, field
from typing import Any, Dict, List, Optional, Tuple

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.formula import ArrayFormula

REQUIRED_SHEETS = ["Total", "Target", "قصادير", "Sales Report", "Data", "Total price", "Statement"]
QASADEER = "قصادير"

# Default category rule: "<amount> SR" or "<amount> ريال" (case/space tolerant).
# Group 1 MUST capture the amount. Editable from the UI.
DEFAULT_CATEGORY_REGEX = r"^(\d+(?:\.\d+)?)\s*(?:sr|s\.r\.?|ريال|ر\.س)$"


class ValidationError(Exception):
    """Raised with a user-readable (Arabic) message when the input cannot be processed."""


# --------------------------------------------------------------------------
# Text normalisation
# --------------------------------------------------------------------------
_DIGITS = str.maketrans("٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹٫", "01234567890123456789.")
_AR_FIX = str.maketrans({"ی": "ي", "ى": "ي", "ک": "ك", "ـ": ""})


def norm_text(value: Any) -> str:
    """Tolerant comparison key: NFKC, Arabic/Persian digits and letter variants, spaces, case."""
    if value is None:
        return ""
    s = unicodedata.normalize("NFKC", str(value)).translate(_DIGITS).translate(_AR_FIX)
    s = re.sub(r"[\u200b-\u200f\u202a-\u202e]", "", s)
    return re.sub(r"\s+", " ", s).strip().casefold()


def _sheet_key(name: str) -> str:
    return "".join(ch for ch in norm_text(name) if ch.isalnum())


_SHEET_KEYS = {_sheet_key(n): n for n in REQUIRED_SHEETS}


def canonical_sheet_name(raw: str) -> Optional[str]:
    """'Target ' -> 'Target', '■ Sales Report\\t' -> 'Sales Report', etc."""
    return _SHEET_KEYS.get(_sheet_key(raw))


def compile_category_regex(pattern: str) -> re.Pattern:
    try:
        rx = re.compile(pattern, re.IGNORECASE)
    except re.error as exc:
        raise ValidationError(f"نمط فئات المبالغ غير صالح: {exc}")
    if rx.groups < 1:
        raise ValidationError("نمط فئات المبالغ يجب أن يحتوي على مجموعة التقاط واحدة على الأقل للمبلغ.")
    return rx


def parse_category(header: Any, rx: Optional[re.Pattern] = None) -> Optional[Tuple[float, str]]:
    """Return (amount, display label) if the header looks like an amount category, else None."""
    if header is None or isinstance(header, bool) or not isinstance(header, (str, int, float)):
        return None
    rx = rx or compile_category_regex(DEFAULT_CATEGORY_REGEX)
    m = rx.match(norm_text(header))
    if not m:
        return None
    try:
        amount = float(m.group(1))
    except (TypeError, ValueError):
        return None
    return amount, f"{amount:g} SR"


# --------------------------------------------------------------------------
# Settings / results
# --------------------------------------------------------------------------
@dataclass
class Settings:
    name_header: str = "اسم العميل"
    number_header: str = "رقم العميل"
    region_header: str = "المدينه"
    category_regex: str = DEFAULT_CATEGORY_REGEX
    # Optional column = sum(face value x count). Off by default: it is a business rule
    # the spec did not define (the legacy sheet uses the same 5/2 weighting).
    include_face_value_total: bool = False


@dataclass
class QasadeerResult:
    table: pd.DataFrame
    categories: List[str]
    source_columns: Dict[str, List[str]]
    issues: List[str] = field(default_factory=list)
    ignored_rows: List[int] = field(default_factory=list)


@dataclass
class Source:
    raw: bytes
    sheet_map: Dict[str, str]  # canonical name -> raw name in the uploaded file
    wb_values: Any
    notes: List[str] = field(default_factory=list)

    def missing_sheets(self) -> List[str]:
        return [n for n in REQUIRED_SHEETS if n not in self.sheet_map]


# --------------------------------------------------------------------------
# 1. Reading
# --------------------------------------------------------------------------
def load_source(file_bytes: bytes) -> Source:
    if not file_bytes:
        raise ValidationError("الملف فارغ.")
    if file_bytes[:2] != b"PK":
        raise ValidationError("صيغة الملف غير مدعومة. احفظ الملف بصيغة xlsx (وليس xlsb أو xls) ثم أعد رفعه.")
    try:
        wb_v = load_workbook(io.BytesIO(file_bytes), data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ValidationError(f"تعذّر فتح ملف Excel: {exc}")
    sheet_map: Dict[str, str] = {}
    for raw in wb_v.sheetnames:
        canon = canonical_sheet_name(raw)
        if canon and canon not in sheet_map:
            sheet_map[canon] = raw
    src = Source(raw=file_bytes, sheet_map=sheet_map, wb_values=wb_v)
    if "Data" not in sheet_map:
        raise ValidationError("ورقة «Data» غير موجودة في الملف، وهي المصدر الأساسي المطلوب.")
    for name in src.missing_sheets():
        src.notes.append(f"الورقة «{name}» غير موجودة في الملف المرفوع، وسيتم إنشاؤها فارغة في الملف الناتج.")
    uncached = _count_uncached_formulas(file_bytes, sheet_map["Data"])
    if uncached:
        src.notes.append(
            f"في ورقة Data حوالي {uncached} خلية تحتوي معادلات بلا قيم محفوظة (الملف حُفظ بأداة لا تحسب المعادلات). "
            "افتح الملف في Excel واحفظه مرة واحدة قبل رفعه، وإلا ستظهر هذه الخلايا فارغة في التقارير."
        )
    return src


def _count_uncached_formulas(file_bytes: bytes, data_sheet: str, max_rows: int = 400) -> int:
    wb_f = load_workbook(io.BytesIO(file_bytes), data_only=False)
    wb_v = load_workbook(io.BytesIO(file_bytes), data_only=True)
    wf, wv = wb_f[data_sheet], wb_v[data_sheet]
    count = 0
    for row_f, row_v in zip(wf.iter_rows(max_row=max_rows), wv.iter_rows(max_row=max_rows)):
        for cf, cv in zip(row_f, row_v):
            v = cf.value
            is_formula = isinstance(v, ArrayFormula) or (isinstance(v, str) and v.startswith("="))
            if is_formula and cv.value is None:
                count += 1
    return count


def read_table(ws) -> Tuple[List[Any], List[List[Any]]]:
    """Header = row 1 (duplicates preserved), rows = everything below."""
    it = ws.iter_rows(values_only=True)
    try:
        headers = list(next(it))
    except StopIteration:
        raise ValidationError("ورقة Data فارغة.")
    return headers, [list(r) for r in it]


def _cell(row: List[Any], idx: Optional[int]) -> Any:
    if idx is None or idx >= len(row):
        return None
    return row[idx]


def _num(value: Any) -> Tuple[Optional[float], bool]:
    """-> (number or None, is_invalid). Blank -> (None, False); text that is not a number -> (None, True)."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None, False
    if isinstance(value, bool):
        return None, True
    if isinstance(value, (int, float)):
        if value != value:  # NaN
            return None, True
        return float(value), False
    if isinstance(value, str):
        s = norm_text(value).replace(",", "")
        try:
            return float(s), False
        except ValueError:
            return None, True
    return None, True


def _find_col(headers: List[Any], wanted: str, before: Optional[int] = None) -> Optional[int]:
    key = norm_text(wanted)
    cands = [i for i, h in enumerate(headers) if norm_text(h) == key]
    if not cands:
        return None
    if before is not None:
        left = [i for i in cands if i < before]
        if left:
            return max(left)  # the one sitting next to the category block
    return cands[0]


def _is_legend(name: str) -> bool:
    return name.strip()[:1] in ("=", "<", ">")


# --------------------------------------------------------------------------
# 2. Computing
# --------------------------------------------------------------------------
def detect_categories(headers: List[Any], category_regex: str = DEFAULT_CATEGORY_REGEX) -> Dict[int, str]:
    rx = compile_category_regex(category_regex)
    out: Dict[int, str] = {}
    for i, h in enumerate(headers):
        parsed = parse_category(h, rx)
        if parsed:
            out[i] = parsed[1]
    return out


def build_qasadeer(headers: List[Any], rows: List[List[Any]], settings: Optional[Settings] = None) -> QasadeerResult:
    s = settings or Settings()
    cats = detect_categories(headers, s.category_regex)
    if not cats:
        raise ValidationError("لم يتم العثور على أي عمود بفئة مبلغ (مثل «5 SR» أو «2 ريال») في ورقة Data.")
    first_cat = min(cats)
    name_idx = _find_col(headers, s.name_header, before=first_cat)
    if name_idx is None:
        raise ValidationError(f"العمود المطلوب «{s.name_header}» غير موجود في ورقة Data. غيّر اسم عمود العميل من الإعدادات.")
    num_idx = _find_col(headers, s.number_header, before=first_cat)
    reg_idx = _find_col(headers, s.region_header, before=first_cat)

    # customer block = rows from the top until two consecutive rows without a valid customer name
    valid_rows: List[int] = []
    blank_run = 0
    for r, row in enumerate(rows):
        nm = _cell(row, name_idx)
        if isinstance(nm, str) and nm.strip() and not _is_legend(nm):
            valid_rows.append(r)
            blank_run = 0
        else:
            blank_run += 1
            if blank_run >= 2 and valid_rows:
                break
    valid_set = set(valid_rows)

    issues: List[str] = []
    order: List[str] = []  # category labels in first-appearance order
    source_columns: Dict[str, List[str]] = {}
    for idx, lab in cats.items():
        if lab not in order:
            order.append(lab)
        source_columns.setdefault(lab, []).append(get_column_letter(idx + 1))

    customers: Dict[str, Dict[str, Any]] = {}
    invalid_cells: List[str] = []
    for r in valid_rows:
        row = rows[r]
        key = norm_text(_cell(row, name_idx))
        c = customers.setdefault(
            key, {"name": re.sub(r"\s+", " ", str(row[name_idx]).strip()), "numbers": [], "region": None, "sums": {}}
        )
        if c["region"] is None and reg_idx is not None and _cell(row, reg_idx) not in (None, ""):
            c["region"] = _cell(row, reg_idx)
        if num_idx is not None:
            n = _cell(row, num_idx)
            if n not in (None, "") and n not in c["numbers"]:
                c["numbers"].append(n)
        for idx, lab in cats.items():
            val, bad = _num(_cell(row, idx))
            if bad:
                invalid_cells.append(f"{get_column_letter(idx + 1)}{r + 2}")
            elif val is not None:
                c["sums"][lab] = c["sums"].get(lab, 0.0) + val

    if invalid_cells:
        issues.append(
            f"تم تجاهل {len(invalid_cells)} خلية غير رقمية في أعمدة الفئات (لم تُحتسب كأصفار)، مثل: "
            + "، ".join(invalid_cells[:8])
        )

    # rows with numbers in category columns that are NOT customer rows (totals, notes, legends)
    ignored: List[int] = []
    for r, row in enumerate(rows):
        if r in valid_set:
            continue
        if any(_num(_cell(row, i))[0] is not None for i in cats):
            ignored.append(r + 2)
    if ignored:
        issues.append(
            f"تم تجاهل {len(ignored)} صفًا خارج كتلة العملاء تحتوي أرقامًا في أعمدة الفئات "
            f"(غالبًا صفوف إجماليات أو ملاحظات)، مثل الصفوف: " + "، ".join(map(str, ignored[:8]))
        )

    multi = [c["name"] for c in customers.values() if len(c["numbers"]) > 1]
    if multi:
        issues.append(f"عملاء بنفس الاسم لهم أكثر من رقم عميل (تم دمجهم باسم العميل): {', '.join(multi[:5])}")
    if num_idx is None:
        issues.append(f"عمود «{s.number_header}» غير موجود؛ التجميع تم حسب اسم العميل فقط.")

    records = []
    for c in customers.values():
        rec: Dict[str, Any] = {}
        if num_idx is not None:
            nums = c["numbers"]
            rec[s.number_header.strip()] = nums[0] if nums else None
        rec[s.name_header.strip()] = c["name"]
        if reg_idx is not None:
            rec["المدينة"] = c["region"]
        for lab in order:
            rec[lab] = c["sums"].get(lab)  # None stays blank: missing is not zero
        if s.include_face_value_total:
            amt = {lab: float(lab.split()[0]) for lab in order}
            rec["اجمالي المبلغ (القيمة × العدد)"] = sum(amt[l] * v for l, v in c["sums"].items())
        records.append(rec)

    df = pd.DataFrame(records)
    for lab in order:
        df[lab] = pd.to_numeric(df[lab], errors="coerce")
    return QasadeerResult(table=df, categories=order, source_columns=source_columns, issues=issues, ignored_rows=ignored)


STATEMENT_HEADERS = [
    "Date", "Invoice No", "Customer - name", "Area", "Amount", "Paying", "المتبقي", "المعلق",
    "PCMO-Min 24L", "PCMO-Min 12L", "PCMO-Syn", "DEO 5L", "DEO 4L", "DEO 1L", "Fluids 4L", "Fluids 12L",
]


def build_statement(
    headers: List[Any],
    rows: List[List[Any]],
    customer: Optional[str] = None,
    search: Optional[str] = None,
    date_from: Optional[dt.date] = None,
    date_to: Optional[dt.date] = None,
) -> pd.DataFrame:
    """Detail rows of the transactions table in Data (no invented mapping: Data's own columns)."""
    idx = {}
    for want in STATEMENT_HEADERS:
        i = _find_col(headers, want)
        if i is not None:
            idx[want] = i
    if "Customer - name" not in idx:
        raise ValidationError("العمود «Customer - name» غير موجود في ورقة Data، لا يمكن إنشاء كشف الحساب.")
    cust_i, date_i = idx["Customer - name"], idx.get("Date")
    want_c = norm_text(customer) if customer else None
    want_s = norm_text(search) if search else None
    out = []
    for row in rows:
        nm = _cell(row, cust_i)
        if not isinstance(nm, str) or not nm.strip():
            continue
        d = _cell(row, date_i) if date_i is not None else None
        d = d.date() if isinstance(d, dt.datetime) else d
        if want_c and norm_text(nm) != want_c:
            continue
        if want_s and want_s not in norm_text(nm):
            continue
        if date_from or date_to:
            if not isinstance(d, dt.date):
                continue
            if date_from and d < date_from:
                continue
            if date_to and d > date_to:
                continue
        out.append({h: (d if h == "Date" else _cell(row, i)) for h, i in idx.items()})
    return pd.DataFrame(out, columns=list(idx))


def target_summary(ws_values) -> Optional[pd.DataFrame]:
    """Monthly Target vs Actual (totals) for 2026 as laid out in the Target sheet; None if layout differs."""
    rows = list(ws_values.iter_rows(min_row=1, max_row=60, values_only=True))
    if len(rows) < 4:
        return None
    start = next((i for i, v in enumerate(rows[0]) if isinstance(v, str) and "2026" in v), None)
    if start is None or start + 5 >= len(rows[0]):
        return None
    total_i, pct_i = start + 4, start + 5
    out, month = [], None
    for r in rows[2:]:
        label = (r[1] or "").strip().lower() if isinstance(r[1], str) else ""
        if isinstance(r[0], str) and r[0].strip():
            month = r[0].strip()
        if not month or label not in ("target", "actual"):
            continue
        if label == "target" and len(out) < 12:
            out.append({"الشهر": month, "المستهدف": r[total_i], "الفعلي": None, "نسبة التحقيق": None})
        elif label == "actual" and out and out[-1]["الشهر"] == month:
            out[-1]["الفعلي"], out[-1]["نسبة التحقيق"] = r[total_i], r[pct_i]
    df = pd.DataFrame(out)
    for col in ("المستهدف", "الفعلي", "نسبة التحقيق"):
        df[col] = pd.to_numeric(df[col], errors="coerce")
    return df if not df.empty else None


# --------------------------------------------------------------------------
# 3. Writing the output workbook
# --------------------------------------------------------------------------
_HDR_FILL = PatternFill("solid", fgColor="1F3864")
_HDR_FONT = Font(bold=True, color="FFFFFF")
_THIN = Side(style="thin", color="BFBFBF")


def _write_df(ws, df: pd.DataFrame, total_row: bool = False) -> None:
    ws.sheet_view.rightToLeft = True
    for j, col in enumerate(df.columns, 1):
        c = ws.cell(row=1, column=j, value=str(col))
        c.fill, c.font = _HDR_FILL, _HDR_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, rec in enumerate(df.itertuples(index=False), 2):
        for j, v in enumerate(rec, 1):
            if v is None or (isinstance(v, float) and v != v) or v is pd.NaT:
                continue
            if isinstance(v, pd.Timestamp):
                v = v.to_pydatetime()
            if hasattr(v, "item") and not isinstance(v, (dt.date, dt.datetime)):
                v = v.item()
            c = ws.cell(row=i, column=j, value=v)
            c.border = Border(bottom=_THIN)
            if isinstance(v, (dt.datetime, dt.date)):
                c.number_format = "yyyy-mm-dd"
            elif isinstance(v, float) and v != int(v):
                c.number_format = "#,##0.00"
    if total_row and len(df):
        last = len(df) + 2
        ws.cell(row=last, column=1, value="الإجمالي").font = Font(bold=True)
        for j, col in enumerate(df.columns, 1):
            if pd.api.types.is_numeric_dtype(df[col]) and str(col) not in ("رقم العميل",):
                c = ws.cell(row=last, column=j, value=float(df[col].sum(skipna=True)))
                c.font = Font(bold=True)
                c.number_format = "#,##0.##"
    ws.freeze_panes = "A2"
    if len(df.columns):
        ws.auto_filter.ref = f"A1:{get_column_letter(len(df.columns))}{max(len(df) + 1, 2)}"
    for j, col in enumerate(df.columns, 1):
        sample = [len(str(x)) for x in df[col].head(200).tolist() if x is not None]
        ws.column_dimensions[get_column_letter(j)].width = min(max([len(str(col))] + sample + [8]) + 3, 40)


def _rename_sheet_with_refs(wb, ws, new_title: str) -> None:
    old = ws.title
    if old == new_title:
        return
    patterns = [f"'{old}'!", f"{old}!"] if re.fullmatch(r"\w+", old) else [f"'{old}'!"]
    repl = f"'{new_title}'!" if not re.fullmatch(r"\w+", new_title) else f"{new_title}!"
    for other in wb.worksheets:
        for row in other.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, ArrayFormula) and isinstance(v.text, str):
                    for p in patterns:
                        v.text = v.text.replace(p, repl)
                elif isinstance(v, str) and v.startswith("="):
                    for p in patterns:
                        v = v.replace(p, repl)
                    c.value = v
    ws.title = new_title


def _statement_is_broken(src: Source) -> bool:
    name = src.sheet_map.get("Statement")
    if not name:
        return False
    for row in src.wb_values[name].iter_rows(values_only=True):
        for v in row:
            if isinstance(v, str) and v in ("#NAME?", "#REF!", "#VALUE!"):
                return True
    return False


def build_output_workbook(
    src: Source,
    qas: QasadeerResult,
    statement_df: Optional[pd.DataFrame] = None,
    force_statement: bool = False,
) -> Tuple[bytes, List[str]]:
    """Keep every original sheet; fix sheet names; regenerate قصادير (and Statement if broken/forced)."""
    notes: List[str] = []
    wb = load_workbook(io.BytesIO(src.raw))  # formulas preserved
    for ws in list(wb.worksheets):
        canon = canonical_sheet_name(ws.title)
        if canon and ws.title != canon:
            _rename_sheet_with_refs(wb, ws, canon)
            notes.append(f"تم تصحيح اسم الورقة «{ws.title.strip()}» إلى «{canon}» مع تحديث المعادلات المرتبطة بها.")

    def replace_sheet(name: str, df: pd.DataFrame, total_row: bool) -> None:
        pos = wb.sheetnames.index(name) if name in wb.sheetnames else len(wb.sheetnames)
        if name in wb.sheetnames:
            wb.remove(wb[name])
        _write_df(wb.create_sheet(name, pos), df, total_row=total_row)

    replace_sheet(QASADEER, qas.table, total_row=True)

    if force_statement or _statement_is_broken(src):
        df = statement_df
        if df is None:
            headers, rows = read_table(src.wb_values[src.sheet_map["Data"]])
            df = build_statement(headers, rows)
        replace_sheet("Statement", df, total_row=False)
        notes.append("تم إنشاء ورقة Statement من بيانات Data لأن معادلاتها الأصلية (FILTER/LAMBDA) لم تكن قابلة للحفظ بشكل صحيح.")

    for name in REQUIRED_SHEETS:
        if name not in wb.sheetnames:
            ws = wb.create_sheet(name)
            ws["A1"] = f"ورقة «{name}» غير موجودة في الملف المصدر."

    wb._sheets = [wb[n] for n in REQUIRED_SHEETS] + [s for s in wb.worksheets if s.title not in REQUIRED_SHEETS]
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), notes
